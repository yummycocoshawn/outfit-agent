"""Excel 讀寫與照片存放。所有使用者共用同一份 xlsx,以『使用者ID』欄區分(暱稱可改、可重名,ID 不會)。"""
import shutil
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core import COLUMNS, clean_text

SEP = "、"
COL_USER, COL_UID, COL_NAME, COL_ITEMS, COL_STYLE = 2, 3, 4, 5, 7


class Store:
    def __init__(self, data_dir):
        self.dir = Path(data_dir)
        self.photos = self.dir / "photos"
        self.photos.mkdir(parents=True, exist_ok=True)
        self.xlsx = self.dir / "outfits.xlsx"
        if not self.xlsx.exists():
            wb = Workbook()
            wb.active.append(COLUMNS)
            wb.save(self.xlsx)

    def add(self, user_id, user_name, rec, photo_path):
        row = [
            rec["time"].strftime("%Y-%m-%d %H:%M:%S"), clean_text(user_name, 40), str(user_id),
            rec["name"], SEP.join(rec["items"]), SEP.join(rec["colors"]), rec["style"],
            "" if rec["temp_min"] is None else f'{rec["temp_min"]}~{rec["temp_max"]}',
            str(photo_path)]
        for attempt in range(5):  # Excel 開著時存檔會被鎖住,等一下再試
            try:
                wb = load_workbook(self.xlsx)
                wb.active.append(row)
                wb.save(self.xlsx)
                self._refresh_view()
                return
            except PermissionError:
                print(f"⚠️ 無法寫入 {self.xlsx}(可能在 Excel 中開啟),第 {attempt + 1} 次重試…")
                time.sleep(2)
        raise PermissionError(f"{self.xlsx} 一直無法寫入")

    def _refresh_view(self):
        """複製一份給人類開啟的檢視檔;使用者開著它時略過,不影響主檔寫入。"""
        try:
            shutil.copyfile(self.xlsx, self.dir / "outfits_view.xlsx")
        except OSError:
            pass

    def records(self, user_id):
        ws = load_workbook(self.xlsx).active
        out = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[COL_UID - 1] != str(user_id) or row[COL_NAME - 1] == "不是穿搭":
                continue
            out.append({"time": datetime.strptime(row[0], "%Y-%m-%d %H:%M:%S"),
                        "name": row[COL_NAME - 1], "items": row[COL_ITEMS - 1].split(SEP),
                        "style": row[COL_STYLE - 1]})
        return out

    def delete_last(self, user_id):
        wb = load_workbook(self.xlsx)
        ws = wb.active
        for idx in range(ws.max_row, 1, -1):
            if ws.cell(idx, COL_UID).value == str(user_id):
                name = ws.cell(idx, COL_NAME).value
                ws.delete_rows(idx)
                wb.save(self.xlsx)
                self._refresh_view()
                return name
        return None
